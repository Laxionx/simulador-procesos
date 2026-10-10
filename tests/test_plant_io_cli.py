"""Lectura del YAML, reportes y línea de comandos."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from procesos import PlantFileError, load_plant
from procesos.cli import main
from procesos.report import element_balance, stream_table, unit_balance, write_csv, write_xlsx

EXAMPLES = Path(__file__).resolve().parents[1] / "ejemplos"

MINIMAL = """
nombre: Mínima
especies: [H2O, {nombre: SiO2, fase: s}]
alimentaciones:
  F: {flujos: {H2O: 10, SiO2: 5}}
equipos:
  - {tipo: divisor, nombre: D, entradas: F, salidas: [A, B], fracciones: [0.4, 0.6]}
"""


def test_carga_los_ejemplos_y_texto_yaml() -> None:
    for path in sorted(EXAMPLES.glob("*.yaml")):
        plant = load_plant(path)
        assert plant.flowsheet.units and plant.name
    plant = load_plant(MINIMAL)
    assert plant.name == "Mínima" and not plant.controllers
    assert plant.flowsheet.solve().streams["A"].flow("H2O") == pytest.approx(4.0)


@pytest.mark.parametrize(
    ("text", "message"),
    [
        ("- solo\n- una lista", "mapa YAML"),
        ("nombre: X\nalimentaciones: {}\nequipos: []", "especies"),
        ("especies: [Xx]\nalimentaciones: {}\nequipos: []", "especies"),
        (MINIMAL.replace("tipo: divisor", "tipo: horno"), "tipo desconocido"),
        (MINIMAL.replace("fracciones: [0.4, 0.6]", "fracciones: [0.4, 0.4]"), "sumar 1"),
        (MINIMAL.replace("{flujos: {H2O: 10, SiO2: 5}}", "{caudal: 3}"), "flujos"),
        (MINIMAL.replace("{flujos: {H2O: 10, SiO2: 5}}", "{flujos: [1, 2]}"), "mapa especie"),
        (MINIMAL.replace("H2O: 10", "H2O: -10"), "negativo"),
        (MINIMAL.replace("entradas: F", "entradas: 3"), "lista de nombres"),
        (MINIMAL.replace("salidas: [A, B]", "salidas: [A, F]"), "iguales a una entrada"),
        (
            MINIMAL + "  - {tipo: mezclador, nombre: E, entradas: [B], salidas: [A]}\n",
            "planta: .*único",
        ),
        (
            MINIMAL + "controladores:\n  - {nombre: C, ajustar: {alimentacion: F}, meta: {}}",
            "falta",
        ),
        (
            MINIMAL
            + "controladores:\n  - {nombre: C, ajustar: {equipo: D, parametro: 'fraccion:A',"
            " minimo: 1, maximo: 0}, meta: {corriente: A, propiedad: total, valor: 1}}",
            "Límites",
        ),
    ],
)
def test_errores_del_archivo_dicen_donde(text: str, message: str) -> None:
    with pytest.raises(PlantFileError, match=message):
        load_plant(text)


def test_reportes_y_balances(tmp_path: Path) -> None:
    plant = load_plant(EXAMPLES / "concentradora_cu.yaml")
    solution = plant.flowsheet.solve()
    rows = stream_table(plant.flowsheet, solution)
    assert rows[0][0] == "Especie (kg/h)" and len(rows) == len(plant.flowsheet.catalog) + 4
    assert all(r.relative_error < 1e-12 for r in element_balance(plant.flowsheet, solution))
    assert all(r.relative_error < 1e-12 for r in unit_balance(plant.flowsheet, solution))
    csv_path = write_csv(rows, tmp_path / "t.csv")
    assert csv_path.read_text(encoding="utf-8-sig").startswith("Especie (kg/h);Mineral")
    sheet = load_workbook(write_xlsx(rows, tmp_path / "t.xlsx")).active
    assert sheet.cell(1, 2).value == "Mineral"


def test_linea_de_comandos(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    code = main(
        ["resolver", str(EXAMPLES / "lixiviacion_oxidos_cu.yaml"), "--salida", str(tmp_path)]
    )
    out = capsys.readouterr().out
    assert code == 0
    assert "Acidez libre" in out and "Peor error de cierre" in out
    assert (tmp_path / "lixiviacion_oxidos_cu.csv").exists()
    assert (tmp_path / "lixiviacion_oxidos_cu.xlsx").exists()

    assert main(["resolver", str(EXAMPLES / "reciclo_simple.yaml")]) == 0
    bad = tmp_path / "mala.yaml"
    bad.write_text("especies: [Xx]\n", encoding="utf-8")
    assert main(["resolver", str(bad)]) == 1
    assert "Error" in capsys.readouterr().err
    assert main(["resolver", str(tmp_path / "no_existe.yaml")]) == 1
